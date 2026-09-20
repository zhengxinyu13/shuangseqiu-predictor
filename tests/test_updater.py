"""``shuangseqiu.updater`` 的测试：解析、双向校验与写盘链路。

全部离线：网络抓取通过注入假抓取函数替换，解析则直接用固定 HTML / JSON 片段。
重点守住「宁可拒绝写盘，也不写进对不上的数据」这条底线。
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import pytest

from shuangseqiu import dataset, updater

TUESDAY = dt.date(2023, 1, 3)
THURSDAY = dt.date(2023, 1, 5)
SUNDAY = dt.date(2023, 1, 8)

# 一部「正常」的 55128 列表页：两行记录，列里故意混入开机号列
HTML_55128 = """
<table>
<tr><th>开奖时间</th><th>期数</th><th>开奖号码</th><th>开机号</th><th>和值</th></tr>
<tr>
  <td>2023-01-08</td>
  <td>2023003</td>
  <td><span class="ball-list red">01</span><span class="ball-list red">07</span>
      <span class="ball-list red">14</span><span class="ball-list red">20</span>
      <span class="ball-list red">27</span><span class="ball-list red">30</span>
      <span class="ball-list kjhblue">10</span></td>
  <td>03,08,17,22,29,30+09</td>
  <td>99</td>
</tr>
<tr>
  <td>2023-01-05</td>
  <td>2023002</td>
  <td><span class="ball-list red">13</span><span class="ball-list red">14</span>
      <span class="ball-list red">15</span><span class="ball-list red">20</span>
      <span class="ball-list red">28</span><span class="ball-list red">33</span>
      <span class="ball-list kjhblue">07</span></td>
  <td></td>
  <td>123</td>
</tr>
<tr><td>没有号码的行</td></tr>
</table>
"""


def _official_payload(entries: list[dict], state: str = "0") -> dict:
    return {"state": state, "message": "查询成功", "result": entries}


def _official_entry(issue: str, red: str, blue: str, day: str, winners: str = "2") -> dict:
    return {
        "code": issue,
        "date": f"{day}(四)",
        "red": red,
        "blue": blue,
        "sales": "335294274",
        "poolmoney": "917243647",
        "prizegrades": [
            {"type": 1, "typenum": winners, "typemoney": "10000000"},
            {"type": 2, "typenum": "95", "typemoney": "276194"},
            {"type": 3, "typenum": "843", "typemoney": "3000"},
        ],
    }


# --------------------------------------------------------------------------
# 解析
# --------------------------------------------------------------------------

def test_parse_55128_reads_balls_date_and_issue() -> None:
    rows = updater.parse_55128(HTML_55128)
    assert [row.issue for row in rows] == [2023003, 2023002]
    assert rows[0].reds == (1, 7, 14, 20, 27, 30)
    assert rows[0].blue == 10
    assert rows[0].date == SUNDAY
    assert rows[1].reds == (13, 14, 15, 20, 28, 33)
    assert rows[1].blue == 7
    assert rows[1].date == THURSDAY


def test_parse_55128_ignores_rows_without_numbers() -> None:
    assert len(updater.parse_55128(HTML_55128)) == 2


def test_parse_55128_raises_when_the_page_has_no_rows() -> None:
    with pytest.raises(updater.UpdateError, match="改版"):
        updater.parse_55128("<html><body>403 Forbidden</body></html>")


def test_parse_55128_is_not_confused_by_the_machine_number_column() -> None:
    """开机号列是后期新增的，按下标取会整体错位；这里按内容特征定位。"""
    rows = updater.parse_55128(HTML_55128)
    assert all(1 <= number <= 33 for row in rows for number in row.reds)
    assert all(1 <= row.blue <= 16 for row in rows)


def test_parse_official_extracts_numbers_and_prize_grades() -> None:
    payload = _official_payload([_official_entry("2023003", "01,07,14,20,27,30", "10", "2023-01-08")])
    rows = updater.parse_official(payload)
    assert len(rows) == 1
    row = rows[0]
    assert row.issue == 2023003
    assert row.reds == (1, 7, 14, 20, 27, 30)
    assert row.blue == 10
    assert row.date == SUNDAY
    assert row.sales == 335294274
    assert row.pool == 917243647
    assert row.first_winners == 2 and row.first_prize == 10000000
    assert row.second_winners == 95 and row.second_prize == 276194


def test_parse_official_tolerates_missing_prize_grades() -> None:
    payload = _official_payload([{"code": "2023003", "date": "2023-01-08(日)",
                                  "red": "01,07,14,20,27,30", "blue": "10"}])
    row = updater.parse_official(payload)[0]
    assert row.sales is None and row.first_winners is None and row.second_prize is None


def test_parse_official_raises_on_failure_state() -> None:
    with pytest.raises(updater.UpdateError, match="官方接口返回失败"):
        updater.parse_official(_official_payload([], state="1"))


def test_parse_official_raises_on_empty_result() -> None:
    with pytest.raises(updater.UpdateError, match="没有返回任何开奖记录"):
        updater.parse_official(_official_payload([]))


def test_remote_draw_converts_to_a_full_record() -> None:
    payload = _official_payload([_official_entry("2023003", "01,07,14,20,27,30", "10", "2023-01-08")])
    record = updater.parse_official(payload)[0].to_record()
    assert isinstance(record, dataset.DrawRecord)
    assert record.issue == 2023003
    assert record.has_bonus is True
    assert record.label == "2023003期"


# --------------------------------------------------------------------------
# 抓取失败要变成可读错误
# --------------------------------------------------------------------------

class _BoomClient:
    """任何请求都抛 httpx 错误的假客户端。"""

    def get(self, *args, **kwargs):
        import httpx

        raise httpx.ConnectError("假装连不上")


def test_fetch_55128_wraps_network_errors() -> None:
    with pytest.raises(updater.UpdateError, match="访问 55128.cn 失败"):
        updater.fetch_55128(client=_BoomClient())


def test_fetch_official_wraps_network_errors() -> None:
    with pytest.raises(updater.UpdateError, match="访问官方接口失败"):
        updater.fetch_official(client=_BoomClient())


class _JsonClient:
    """返回固定 JSON 的假客户端。"""

    def __init__(self, payload: dict) -> None:
        self.payload = payload

    def get(self, *args, **kwargs):
        class _Response:
            def __init__(self, payload: dict) -> None:
                self._payload = payload

            def raise_for_status(self) -> None:
                return None

            def json(self) -> dict:
                return self._payload

        return _Response(self.payload)


def test_fetch_official_wraps_bad_json() -> None:
    class _Bad(_JsonClient):
        def get(self, *args, **kwargs):
            class _Response:
                def raise_for_status(self) -> None:
                    return None

                def json(self):
                    raise ValueError("不是 JSON")

            return _Response()

    with pytest.raises(updater.UpdateError, match="不是合法 JSON"):
        updater.fetch_official(client=_Bad({}))


# --------------------------------------------------------------------------
# 双向校验
# --------------------------------------------------------------------------

@pytest.fixture
def local_records(make_record):
    """三期本地数据，最新 2023003。"""

    def _build():
        return [
            make_record(2023001, (2, 6, 9, 15, 22, 31), 4, day=TUESDAY),
            make_record(2023002, (13, 14, 15, 20, 28, 33), 7, day=THURSDAY),
            make_record(2023003, (1, 7, 14, 20, 27, 30), 10, day=SUNDAY),
        ]

    return _build()


def _sources(primary, official):
    return {updater.SOURCE_PRIMARY: primary, updater.SOURCE_OFFICIAL: official}


def test_plan_without_new_issues_is_a_no_op(local_records, make_remote) -> None:
    same = [make_remote(record.issue, record.reds, record.blue, record.date) for record in local_records]
    plan = updater.build_update_plan(local_records, _sources(same, same))
    assert plan.has_new is False
    assert plan.conflicts == ()
    assert plan.can_write is True
    assert plan.local_latest == 2023003
    assert plan.remote_latest == 2023003


def test_plan_accepts_a_new_issue_when_both_sources_agree(local_records, make_remote) -> None:
    fresh = make_remote(2023004, (5, 9, 12, 23, 26, 32), 15, dt.date(2023, 1, 10),
                        sales=300000000, first_winners=3)
    plan = updater.build_update_plan(
        local_records, _sources(list(local_records) + [fresh], [fresh])
    )
    assert plan.has_new is True
    assert plan.can_write is True
    assert [record.issue for record in plan.new_records] == [2023004]
    assert plan.new_records[0].sales == 300000000
    assert plan.remote_latest == 2023004


def test_plan_refuses_a_new_issue_missing_from_55128(local_records, make_remote) -> None:
    fresh = make_remote(2023004, (5, 9, 12, 23, 26, 32), 15, dt.date(2023, 1, 10))
    plan = updater.build_update_plan(local_records, _sources(list(local_records), [fresh]))
    assert plan.has_new is False
    assert plan.can_write is False
    assert "2023004期" in plan.conflicts[0]
    assert "55128.cn" in plan.conflicts[0]


def test_plan_refuses_a_new_issue_missing_from_official(local_records, make_remote) -> None:
    fresh = make_remote(2023004, (5, 9, 12, 23, 26, 32), 15, dt.date(2023, 1, 10))
    plan = updater.build_update_plan(local_records, _sources([fresh], list(local_records)))
    assert plan.can_write is False
    assert "官方" in plan.conflicts[0]


@pytest.mark.parametrize(
    ("mutation", "expected"),
    [
        ({"reds": (5, 9, 12, 23, 26, 33)}, "红球不同"),
        ({"blue": 16}, "蓝球不同"),
        ({"day": dt.date(2023, 1, 12)}, "开奖日期不同"),
    ],
)
def test_plan_refuses_when_the_two_sources_disagree(local_records, make_remote, mutation, expected) -> None:
    base = dict(reds=(5, 9, 12, 23, 26, 32), blue=15, day=dt.date(2023, 1, 10))
    agreed = make_remote(2023004, **base)
    mutated = make_remote(2023004, **{**base, **mutation})
    plan = updater.build_update_plan(
        local_records, _sources(list(local_records) + [agreed], [mutated])
    )
    assert plan.can_write is False
    assert expected in plan.conflicts[0]
    assert "2023004期" in plan.conflicts[0]


def test_plan_detects_a_source_tampering_with_history(local_records, make_remote) -> None:
    """重叠期回头比对：某个源偷偷改了老期号也必须被抓住。"""
    tampered = [
        make_remote(2023001, (2, 6, 9, 15, 22, 31), 4, TUESDAY),
        make_remote(2023002, (13, 14, 15, 20, 28, 99), 7, THURSDAY),  # 号码被改
        make_remote(2023003, (1, 7, 14, 20, 27, 30), 10, SUNDAY),
    ]
    plan = updater.build_update_plan(local_records, _sources(tampered, list(local_records)))
    assert plan.can_write is False
    assert any("2023002期" in conflict for conflict in plan.conflicts)
    assert any("历史数据与本地库冲突" in conflict for conflict in plan.conflicts)


def test_plan_records_cross_check_rows(local_records, make_remote) -> None:
    same = [make_remote(record.issue, record.reds, record.blue, record.date) for record in local_records]
    plan = updater.build_update_plan(local_records, _sources(same, same))
    categories = [row[0] for row in plan.check_rows]
    assert "数据来源" in categories and "交叉校验" in categories and "更新" in categories
    overlap = [row for row in plan.check_rows if "重叠期" in row[1]]
    assert len(overlap) == 2
    assert all(row[2] == "通过" for row in overlap)
    assert {name: periods for name, periods, _ in plan.source_rows}[updater.SOURCE_PRIMARY] == 3


def test_plan_requires_both_sources(local_records) -> None:
    with pytest.raises(ValueError, match="缺少数据源"):
        updater.build_update_plan(local_records, {updater.SOURCE_PRIMARY: []})


# --------------------------------------------------------------------------
# 写盘链路（全部在临时工作簿上跑，绝不碰真实文件）
# --------------------------------------------------------------------------

def test_apply_update_writes_the_new_period(mini_workbook, local_records, make_remote) -> None:
    path = mini_workbook(local_records)
    fresh = make_remote(2023004, (5, 9, 12, 23, 26, 32), 15, dt.date(2023, 1, 10),
                        sales=300000000, pool=700000000, first_winners=3,
                        first_prize=8888888, second_winners=77, second_prize=222222)
    plan = updater.build_update_plan(local_records, _sources(list(local_records) + [fresh], [fresh]))
    written, total = updater.apply_update(path, plan)

    assert total == 4
    reloaded = dataset.read_records(written)
    assert [record.issue for record in reloaded] == [2023001, 2023002, 2023003, 2023004]
    latest = reloaded[-1]
    assert latest.reds == (5, 9, 12, 23, 26, 32)
    assert latest.blue == 15
    assert latest.sales == 300000000 and latest.first_prize == 8888888

    # 6 张表都要在，且派生列跟着重算
    import openpyxl

    workbook = openpyxl.load_workbook(written, read_only=True, data_only=True)
    try:
        assert tuple(workbook.sheetnames) == dataset.SHEET_ORDER
        assert workbook[dataset.SHEET_RECORDS].max_row == 5
        assert workbook[dataset.SHEET_BONUS].max_row == 5
        assert workbook[dataset.SHEET_RED_STATS].max_row == 34
        rows = list(workbook[dataset.SHEET_RECORDS].iter_rows(min_row=2, max_row=2, values_only=True))
        assert rows[0][0] == 2023004, "最新一期必须排在「开奖记录」最上面"
        assert rows[0][12] == sum((5, 9, 12, 23, 26, 32)), "和值没重算"
    finally:
        workbook.close()


def test_apply_update_refuses_a_plan_with_conflicts(mini_workbook, local_records, make_remote) -> None:
    path = mini_workbook(local_records)
    fresh = make_remote(2023004, (5, 9, 12, 23, 26, 32), 15, dt.date(2023, 1, 10))
    plan = updater.build_update_plan(local_records, _sources([fresh], list(local_records)))
    before = path.read_bytes()
    with pytest.raises(ValueError, match="不能写盘"):
        updater.apply_update(path, plan)
    assert path.read_bytes() == before, "有冲突时文件不能被改动"


def test_check_and_update_writes_when_a_new_issue_is_confirmed(
    mini_workbook, local_records, make_remote
) -> None:
    path = mini_workbook(local_records)
    fresh = make_remote(2023004, (5, 9, 12, 23, 26, 32), 15, dt.date(2023, 1, 10), sales=999)
    fetchers = {
        updater.SOURCE_PRIMARY: lambda: list(local_records) + [fresh],
        updater.SOURCE_OFFICIAL: lambda: [fresh],
    }
    result = updater.check_and_update(path, fetchers)
    assert result.status == "updated" and result.ok
    assert result.total_records == 4
    assert "2023004期" in result.message
    assert dataset.read_records(path)[-1].issue == 2023004


def test_check_and_update_reports_up_to_date(mini_workbook, local_records, make_remote) -> None:
    path = mini_workbook(local_records)
    same = [make_remote(record.issue, record.reds, record.blue, record.date) for record in local_records]
    result = updater.check_and_update(
        path, {updater.SOURCE_PRIMARY: lambda: same, updater.SOURCE_OFFICIAL: lambda: same}
    )
    assert result.status == "current" and result.ok
    assert "已是最新" in result.message
    assert result.total_records == 3


def test_check_and_update_refuses_and_says_which_issue_disagrees(
    mini_workbook, local_records, make_remote
) -> None:
    path = mini_workbook(local_records)
    fresh = make_remote(2023004, (5, 9, 12, 23, 26, 32), 15, dt.date(2023, 1, 10))
    wrong = make_remote(2023004, (5, 9, 12, 23, 26, 33), 15, dt.date(2023, 1, 10))
    before = path.read_bytes()
    result = updater.check_and_update(
        path,
        {updater.SOURCE_PRIMARY: lambda: list(local_records) + [fresh],
         updater.SOURCE_OFFICIAL: lambda: [wrong]},
    )
    assert result.status == "conflict" and not result.ok
    assert "2023004期" in result.message and "红球不同" in result.message
    assert path.read_bytes() == before, "冲突时绝不能写盘"


def test_check_and_update_turns_a_fetch_failure_into_a_message(
    mini_workbook, local_records
) -> None:
    path = mini_workbook(local_records)

    def boom():
        raise updater.UpdateError("假装站点挂了")

    result = updater.check_and_update(
        path, {updater.SOURCE_PRIMARY: boom, updater.SOURCE_OFFICIAL: lambda: []}
    )
    assert result.status == "error"
    assert "55128.cn" in result.message and "假装站点挂了" in result.message


def test_check_and_update_reports_an_unreadable_data_file(tmp_path: Path) -> None:
    result = updater.check_and_update(
        tmp_path / "不存在.xlsx",
        {updater.SOURCE_PRIMARY: lambda: [], updater.SOURCE_OFFICIAL: lambda: []},
    )
    assert result.status == "error"
    assert "读取本地数据失败" in result.message


def test_check_and_update_requires_every_source(mini_workbook, local_records) -> None:
    path = mini_workbook(local_records)
    result = updater.check_and_update(path, {updater.SOURCE_PRIMARY: lambda: []})
    assert result.status == "error"
    assert "未提供数据源" in result.message


def test_updates_are_applied_on_top_of_each_other(mini_workbook, local_records, make_remote) -> None:
    """连着更新两次（第二次带两期）也要正确累积。"""
    path = mini_workbook(local_records)
    first = make_remote(2023004, (5, 9, 12, 23, 26, 32), 15, dt.date(2023, 1, 10), sales=1)
    updater.check_and_update(
        path,
        {updater.SOURCE_PRIMARY: lambda: list(local_records) + [first],
         updater.SOURCE_OFFICIAL: lambda: [first]},
    )
    second = make_remote(2023005, (3, 8, 19, 24, 29, 33), 2, dt.date(2023, 1, 12), sales=2)
    result = updater.check_and_update(
        path,
        {updater.SOURCE_PRIMARY: lambda: [first, second],
         updater.SOURCE_OFFICIAL: lambda: [first, second]},
    )
    assert result.status == "updated"
    assert [record.issue for record in dataset.read_records(path)] == [
        2023001, 2023002, 2023003, 2023004, 2023005
    ]
