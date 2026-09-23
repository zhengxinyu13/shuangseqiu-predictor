"""双色球历史开奖数据的定位、读取与解析。

数据契约（由 ``tests/test_data.py`` 负责守住）：

- 工作簿为 ``双色球历史开奖数据_全量.xlsx``，含 6 个工作表
  （开奖记录 / 红球统计 / 蓝球统计 / 年度概况 / 奖金与奖池 / 数据校验）；
- 开奖明细在 ``开奖记录`` 工作表，表头 18 列，见 :data:`HEADER`；
- 红球拆成 ``红球1`` ~ ``红球6`` 六列存放（已按升序），蓝球单列 ``蓝球``；
- 期号为 7 位整数（如 ``2026108``）= 4 位年份 + 3 位年内序号，
  对外统一成 ``2026108期`` 标签形式，便于报告直接展示；
- 行序为倒序，最新一期在最前。
"""

from __future__ import annotations

import re
import shutil
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Sequence

import openpyxl


def application_root() -> Path:
    """程序根目录：源码运行时是仓库根，打包成 exe 后是 exe 所在目录。

    这一步对打包分发是**必需的**。PyInstaller 会把模块解包到临时目录，
    此时 ``__file__`` 指向那个临时目录，退出即删 —— 如果数据目录还按
    ``__file__`` 推算，「检查更新」写进去的新期号一关窗口就没了。
    所以冻结运行时必须改用 ``sys.executable``（exe 自己的路径）来定位。
    """
    if getattr(sys, "frozen", False):  # PyInstaller / cx_Freeze 打包后为真
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent.parent


PROJECT_ROOT = application_root()
DEFAULT_DATA_DIR = PROJECT_ROOT / "data"
DATA_FILE_NAME = "双色球历史开奖数据_全量.xlsx"


def bundled_resource_dir() -> Path | None:
    """打包时被塞进 exe 内部的只读资源目录；源码运行时返回 ``None``。

    PyInstaller 解包后会把路径写在 ``sys._MEIPASS``。
    """
    base = getattr(sys, "_MEIPASS", None)
    return Path(base) if base else None


def seed_default_data_file() -> Path:
    """返回应当使用的数据文件路径；首次运行时先从 exe 内部释放一份初始数据。

    分发的 exe 里带着一份完整的历史数据作为「种子」（打包时的快照）。
    第一次在别人电脑上运行时，把它复制到 **exe 旁边** 的 ``data/`` 里，
    之后「检查更新」写的就是这个可写、可备份、可单独替换的副本，
    而不是 exe 内部的只读资源。

    已存在则原样返回，绝不覆盖 —— 用户自己更新过的数据不能被上线时的旧种子冲掉。
    """
    target = DEFAULT_DATA_DIR / DATA_FILE_NAME
    if target.is_file():
        return target

    bundled = bundled_resource_dir()
    if bundled is None:  # 源码运行：保持原样，交给调用方报「找不到数据文件」
        return target

    source = bundled / "data" / DATA_FILE_NAME
    if source.is_file():
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, target)
    return target


SHEET_NAME = "开奖记录"
HEADER: tuple[str, ...] = (
    "期号",
    "年份",
    "期序",
    "开奖日期",
    "星期",
    "红球1",
    "红球2",
    "红球3",
    "红球4",
    "红球5",
    "红球6",
    "蓝球",
    "和值",
    "跨度",
    "奇偶比",
    "大小比",
    "三区比",
    "连号",
)

RED_BALL_COUNT = 6
RED_BALL_MIN = 1
RED_BALL_MAX = 33
BLUE_BALL_MIN = 1
BLUE_BALL_MAX = 16

# 列位置从表头派生而不是写死：表头一调整，索引自动跟随，不会错位。
_COLUMN_INDEX = {name: position for position, name in enumerate(HEADER)}
ISSUE_COLUMN = _COLUMN_INDEX["期号"]
RED_COLUMNS: tuple[int, ...] = tuple(_COLUMN_INDEX[f"红球{n}"] for n in range(1, RED_BALL_COUNT + 1))
BLUE_COLUMN = _COLUMN_INDEX["蓝球"]

# 期号允许两种写法：数据源里的纯数字（2026108）与展示用的标签（2026108期）
_ISSUE_PATTERN = re.compile(r"^(\d{4})(\d{3})(?:期)?$")
_BALLS_PATTERN = re.compile(r"^\d{2}(?: \d{2}){6}$")


@dataclass(frozen=True)
class Draw:
    """一期开奖记录。"""

    issue: str
    year: int
    index: int
    reds: tuple[int, ...]
    blue: int

    @property
    def sort_key(self) -> tuple[int, int]:
        """可比较的期次键，越大表示越新。"""
        return (self.year, self.index)


def normalize_issue(value: Any) -> str:
    """把期号统一成 ``2026108期`` 标签形式。

    Args:
        value: 期号原值，可能是整数（Excel 单元格）、纯数字文本或已带「期」的标签。

    Returns:
        形如 ``2026108期`` 的字符串。

    Raises:
        ValueError: 期号为空或类型不受支持时。
    """
    if isinstance(value, bool) or not isinstance(value, (int, str)):
        raise ValueError(f"期号类型非法：{value!r}，期望整数或文本")
    text = str(value).strip()
    if text.endswith(".0"):  # Excel 偶尔把整数读成浮点文本
        text = text[:-2]
    if not text:
        raise ValueError("期号为空")
    return text if text.endswith("期") else f"{text}期"


def parse_issue(text: Any) -> tuple[int, int]:
    """把 ``2026108期`` / ``2026108`` 解析成 ``(2026, 108)``。

    Raises:
        ValueError: 期号不是「4 位年份 + 3 位序号」，或多余字符时。
    """
    matched = _ISSUE_PATTERN.match(str(text).strip())
    if matched is None:
        raise ValueError(f"期号格式非法：{text!r}，期望形如 '2026108期' 或 '2026108'")
    return int(matched.group(1)), int(matched.group(2))


def parse_balls(text: str) -> tuple[tuple[int, ...], int]:
    """把 ``06 11 13 14 20 28 16`` 解析成 ``((6, 11, 13, 14, 20, 28), 16)``。

    Raises:
        ValueError: 号码不是 7 个两位数字、以单个空格分隔时。
    """
    normalized = str(text).strip()
    if _BALLS_PATTERN.match(normalized) is None:
        raise ValueError(f"号码格式非法：{text!r}，期望 7 个两位数字、以单个空格分隔")
    numbers = [int(part) for part in normalized.split(" ")]
    return tuple(numbers[:RED_BALL_COUNT]), numbers[RED_BALL_COUNT]


def parse_draw(issue: str, balls: str) -> Draw:
    """把期号与号码串解析成 :class:`Draw`。

    Raises:
        ValueError: 期号或号码格式非法时。
    """
    year, index = parse_issue(issue)
    reds, blue = parse_balls(balls)
    return Draw(issue=normalize_issue(issue), year=year, index=index, reds=reds, blue=blue)


def parse_row(row: Sequence[Any]) -> Draw:
    """把「开奖记录」工作表的一行解析成 :class:`Draw`。

    红球以 6 个独立整数列存放，这里先归一成 ``06 11 13 14 20 28 16`` 标准号码串，
    再交给 :func:`parse_draw` —— 让号码校验只存在于一处，避免出现第二套解析逻辑。

    Args:
        row: 一行的单元格值序列，长度至少到 ``蓝球`` 列。

    Returns:
        解析后的 :class:`Draw`。

    Raises:
        ValueError: 行长度不足、期号非法或红球/蓝球含非整数值时。
    """
    if len(row) <= BLUE_COLUMN:
        raise ValueError(f"行长度不足：仅 {len(row)} 列，至少需要 {BLUE_COLUMN + 1} 列")
    issue = normalize_issue(row[ISSUE_COLUMN])
    numbers = [row[position] for position in RED_COLUMNS] + [row[BLUE_COLUMN]]
    try:
        balls = " ".join(f"{int(value):02d}" for value in numbers)
    except (TypeError, ValueError) as error:
        raise ValueError(f"期号 {issue} 的红球/蓝球存在非整数值：{numbers!r}") from error
    return parse_draw(issue, balls)


def find_data_file(data_dir: Path | str | None = None) -> Path:
    """定位历史开奖数据文件。

    Raises:
        FileNotFoundError: 文件不存在时。
    """
    directory = Path(data_dir) if data_dir is not None else DEFAULT_DATA_DIR
    path = directory / DATA_FILE_NAME
    if not path.is_file():
        raise FileNotFoundError(f"数据文件不存在：{path}")
    return path


def load_rows(path: Path | str | None = None) -> list[tuple]:
    """读取 ``开奖记录`` 工作表全部行（含表头），返回元组列表。

    Raises:
        ValueError: 工作簿缺少 ``开奖记录`` 工作表时。
    """
    target = Path(path) if path is not None else find_data_file()
    workbook = openpyxl.load_workbook(target, read_only=True, data_only=True)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ValueError(f"工作表 {SHEET_NAME!r} 不存在，实际为 {workbook.sheetnames!r}")
        worksheet = workbook[SHEET_NAME]
        return [tuple(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def load_draws(path: Path | str | None = None) -> list[Draw]:
    """读取并解析全部开奖记录，顺序与文件一致（最新在前）。

    Raises:
        ValueError: 文件完全为空，或表头与 :data:`HEADER` 不符时。
    """
    rows = load_rows(path)
    if not rows:
        raise ValueError("数据文件为空，连表头都没有")
    header = tuple("" if cell is None else str(cell) for cell in rows[0])
    if header != HEADER:
        raise ValueError(f"表头不符合预期：{header!r} != {HEADER!r}")
    return [parse_row(row) for row in rows[1:]]
